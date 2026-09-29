#!/usr/bin/env python3
"""Build the testing workbook: what to check in game, what is shelved, what the listing needs.

Usage: python3 tools/testing_xlsx.py [--out ~/Desktop/ForeverAddons/ForeverBuddy-testing.xlsx]

Three sheets, all meant to be filled in over time: a checklist with a status per build, the
things switched off and what would bring each one back, and the CurseForge steps.
"""
import argparse
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
BAND = PatternFill("solid", fgColor="EDF2FA")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
STATUSES = ["Not tested", "Works", "Broken", "Shelved"]

# area, what to check, how, what should happen, status as at 0.20.2
CHECKS = [
    ("The window", "Welcome page", "/fb", "Five sections, no tally under Made by Leemo", "Works"),
    ("The window", "Rail icons", "Click each icon down the left", "All five open their page", "Works"),
    ("The window", "Headings jump", "Click a heading on the welcome page", "Opens that screen", "Works"),
    ("The window", "Escape and dragging", "Escape, then drag the title bar", "Closes, and moves", "Works"),
    ("Dungeon journal", "Covers", "Open the journal", "22 dungeons, loading screen behind each", "Works"),
    ("Dungeon journal", "Counts and emblems", "Look at the top right of a card", "Quest count, emblem where one side only", "Works"),
    ("Dungeon journal", "Chains", "Click a dungeon", "Each chain on its own panel", "Works"),
    ("Dungeon journal", "Done and in-log marks", "Open one you have quests in", "Tick on done, question mark on in your log", "Works"),
    ("Dungeon journal", "Step colours", "Read the line under a step", "Grey to pick up, orange in log, green done", "Works"),
    ("Dungeon journal", "Faction down a chain", "Open the Deadmines", "Blue from the first step to the last", "Works"),
    ("Dungeon journal", "Arrow", "Click a step", "Arrow points at the giver", "Works"),
    ("Dungeon journal", "Only quests I can take", "Toggle it on the grid and inside a dungeon", "Counts and trees follow it", "Works"),
    ("Dungeon journal", "Reward on a tooltip", "Hover a step with a known reward", "Names the item", "Works"),
    ("Abilities", "Your own class", "/fb abilities", "What you can train now, then every level to 60", "Works"),
    ("Abilities", "Spell tooltip", "Hover a spell", "Real tooltip, plus what a class trainer charges", "Works"),
    ("Abilities", "Known spells", "Look at the ready list", "Nothing you already know", "Works"),
    ("Abilities", "Another class", "/fb abilities mage", "That class's ladder", "Works"),
    ("Abilities", "Race-only spells", "Look at a priest at level 10", "Race named, and only yours shown", "Works"),
    ("Where to level", "The list", "/fb zones", "Zones by level per continent", "Works"),
    ("Where to level", "Click a zone", "Click a row", "Map opens at that continent", "Works"),
    ("Where to level", "Lighting a zone up", "Click a row", "SHELVED: the art greys the map", "Shelved"),
    ("Where to level", "Ranges on the map", "Open a continent map", "A range over each zone, in the right place", "Works"),
    ("City maps", "Pins appear", "Open a capital", "Pins, and a panel top left", "Works"),
    ("City maps", "The six switches", "Tick each one off and on", "Its pins come and go", "Works"),
    ("City maps", "Per class and profession", "Expand either trainer row", "A tick each, with All and None", "Works"),
    ("City maps", "Icons", "Look at the pins", "Class icons and trade icons", "Works"),
    ("City maps", "Weapon masters", "Look for the sword pins", "One or two per city", "Works"),
    ("City maps", "Stormwind", "Open Stormwind", "Pins in the right places, harbour and all", "Works"),
    ("City maps", "Switches remembered", "Turn some off, log out and back", "Still off", "Works"),
    ("Tooltips", "Item uses", "Hover any item", "Quests and recipes that need it, keep or vendor", "Works"),
    ("Tooltips", "Item scores", "Hover a piece of gear", "SHELVED: Forever's spell power wording", "Shelved"),
    ("Tooltips", "Delete warnings", "Destroy a quest item", "SHELVED: warning never appears", "Shelved"),
    ("At the vendor", "Repair and sell junk", "Open a vendor", "Repairs, sells greys, one line each", "Works"),
    ("At the vendor", "Equip prompts", "Loot an upgrade", "SHELVED: judged by the broken scores", "Shelved"),
    ("Quests", "Automate quests", "Turn it on, talk to a giver", "Accepts and hands in, Shift pauses", "Works"),
    ("Everything else", "Hide Lua errors", "Leave it ticked", "No error popups", "Works"),
    ("Everything else", "Bag icon markers", "Open your bags", "SHELVED: markers do not paint", "Shelved"),
    ("Everything else", "Trainer recording", "Open any trainer", "Says how many spells it recorded", "Works"),
    ("Everything else", "Saved file loads", "/fb collect check", "Says the saved file loaded at login", "Works"),
    ("Everything else", "Update notice", "Group with someone on an older build", "They are told, you are not", "Not tested"),
]

SHELVED = [
    ("Item scores on tooltips", "Gear/Score.lua",
     'Forever writes spell power as "increases damage and healing done by magical spells and effects by up to N"',
     "Read that wording in the stat parser", "Brings back equip prompts and the bag upgrade arrow too"),
    ("Equip upgrade prompts", "Automation/AutoEquip.lua", "Judges an upgrade by the item scores",
     "Fix the scores above", ""),
    ("Delete warnings", "Tooltips/DeleteWarn.lua", "The warning never appears on Forever's confirmation",
     "Find which popup Forever really uses", ""),
    ("Bag icon markers", "Bags/Marks.lua", "The markers do not paint on Forever's bag frames",
     "Find the frame the client really draws", ""),
    ("Lighting a zone on the map", "Nav/MapLevels.lua", "The client's highlight art greys the whole map",
     "Draw the zone's own shape, or a neater outline", "Clicking a zone still opens the map"),
]

RELEASE = [
    ("Decide", "Licence", "GPL-3.0, crediting Questie", "Done"),
    ("Decide", "Public repo", "Fresh repo from a scrubbed tree, this one kept private", "Not started"),
    ("Make", "Logo", "400 by 400, your own artwork, no Blizzard images", "Not started"),
    ("Make", "Summary", "One sentence", "Not started"),
    ("Make", "Description", "What it does, the commands, what the recorder stores. No download links", "Not started"),
    ("Make", "Screenshots", "Journal, city map, abilities", "Not started"),
    ("Upload", "Create the project", "CurseForge, two-factor on, World of Warcraft, Addons", "Not started"),
    ("Upload", "First file", "Upload the zip as a Release, Forever game version", "Not started"),
    ("Upload", "Moderation", "A person reviews the first file, usually same day", "Not started"),
    ("After", "Project id", "Add X-Curse-Project-ID to the TOC", "Not started"),
    ("After", "API token", "CF_API_KEY as a repository secret", "Not started"),
    ("After", "Tag to release", "git tag -a v1.0.0 && git push origin v1.0.0", "Not started"),
]


def head(ws, row, names, widths):
    for i, (name, width) in enumerate(zip(names, widths), start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, italic=True, size=9)


def write(ws, row, values, fill_from=None):
    for i, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=value)
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(vertical="top", wrap_text=i > 2)
        if fill_from and i >= fill_from:
            cell.fill = FILL_IN
        elif row % 2 == 0:
            cell.fill = BAND


def status_rule(ws, column, first, last):
    check = DataValidation(type="list", formula1='"%s"' % ",".join(STATUSES), allow_blank=True)
    ws.add_data_validation(check)
    check.add("%s%d:%s%d" % (column, first, column, last))
    span = "%s%d:%s%d" % (column, first, column, last)
    ws.conditional_formatting.add(span, CellIsRule(operator="equal", formula=['"Works"'],
                                                   fill=PatternFill("solid", fgColor="D8F0D8")))
    ws.conditional_formatting.add(span, CellIsRule(operator="equal", formula=['"Broken"'],
                                                   fill=PatternFill("solid", fgColor="F8D7D7")))
    ws.conditional_formatting.add(span, CellIsRule(operator="equal", formula=['"Shelved"'],
                                                   fill=PatternFill("solid", fgColor="EDE4CF")))


def checklist(wb):
    ws = wb.active
    ws.title = "Test checklist"
    title(ws, "ForeverBuddy: what to check in game",
          "Status as at 0.20.2. Fill in the build you tested, the date and anything that went wrong. "
          "Shelved means it is switched off on purpose; see the Shelved sheet.")
    head(ws, 4, ["Area", "What", "How", "What should happen", "Status", "Build", "Date", "Notes"],
         [18, 24, 30, 46, 13, 10, 12, 40])
    for i, (area, what, how, expect, status) in enumerate(CHECKS, start=5):
        write(ws, i, [area, what, how, expect, status, "0.20.2" if status != "Not tested" else "", "", ""],
              fill_from=6)
    last = len(CHECKS) + 4
    status_rule(ws, "E", 5, last)
    ws.cell(row=last + 2, column=1, value="Working:").font = Font(name=FONT, bold=True)
    ws.cell(row=last + 2, column=2, value='=COUNTIF(E5:E%d,"Works")' % last).font = Font(name=FONT)
    ws.cell(row=last + 3, column=1, value="Broken:").font = Font(name=FONT, bold=True)
    ws.cell(row=last + 3, column=2, value='=COUNTIF(E5:E%d,"Broken")' % last).font = Font(name=FONT)
    ws.cell(row=last + 4, column=1, value="Shelved:").font = Font(name=FONT, bold=True)
    ws.cell(row=last + 4, column=2, value='=COUNTIF(E5:E%d,"Shelved")' % last).font = Font(name=FONT)
    ws.cell(row=last + 5, column=1, value="Still to test:").font = Font(name=FONT, bold=True)
    ws.cell(row=last + 5, column=2, value='=COUNTIF(E5:E%d,"Not tested")' % last).font = Font(name=FONT)


def shelved(wb):
    ws = wb.create_sheet("Shelved")
    title(ws, "Switched off for the first release",
          "Each is a SHIPPING constant at the top of its file. The code and its tests are untouched.")
    head(ws, 4, ["What", "File", "Why", "What would bring it back", "Knock-on", "Fixed in"],
         [28, 28, 54, 40, 40, 12])
    for i, row in enumerate(SHELVED, start=5):
        write(ws, i, list(row) + [""], fill_from=6)


def release(wb):
    ws = wb.create_sheet("Release")
    title(ws, "Getting it onto CurseForge", "In order. Nothing below Upload can start until the logo exists.")
    head(ws, 4, ["Stage", "Step", "What it means", "Status", "Notes"], [12, 20, 62, 14, 40])
    for i, row in enumerate(RELEASE, start=5):
        write(ws, i, list(row) + [""], fill_from=5)
    check = DataValidation(type="list", formula1='"Not started,In progress,Done"', allow_blank=True)
    ws.add_data_validation(check)
    check.add("D5:D%d" % (len(RELEASE) + 4))
    ws.conditional_formatting.add("D5:D%d" % (len(RELEASE) + 4),
                                  CellIsRule(operator="equal", formula=['"Done"'],
                                             fill=PatternFill("solid", fgColor="D8F0D8")))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=os.path.expanduser("~/Desktop/ForeverAddons/ForeverBuddy-testing.xlsx"))
    args = parser.parse_args(argv)
    wb = Workbook()
    checklist(wb)
    shelved(wb)
    release(wb)
    wb.save(args.out)
    print("%d checks, %d shelved, %d release steps -> %s" % (len(CHECKS), len(SHELVED), len(RELEASE), args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
