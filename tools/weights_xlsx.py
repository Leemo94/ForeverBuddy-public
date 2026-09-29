#!/usr/bin/env python3
"""Rebuild the stat weights workbook from Data/Weights.lua, the numbers the addon really scores with.

Usage: python3 tools/weights_xlsx.py [--out ~/Desktop/ForeverAddons/ForeverBuddy-stat-weights.xlsx]

Every spec links to the WoWSims file its numbers came from, pinned to the commit they were read
at, so any figure can be checked at source.
"""
import argparse
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SPEC = re.compile(
    r'\["([A-Za-z]+)"\] = \{ class = "([A-Z]+)", name = "([^"]+)", source = "([^"]*)",'
    r' stats = \{([^}]*)\}, pseudo = \{([^}]*)\} \}')
INFO = re.compile(r'ns\.WeightsInfo = \{ game = "([^"]+)", generated = "([^"]+)", specs = (\d+),'
                  r' repo = "([^"]+)", commit = "([^"]+)" \}')
PAIR = re.compile(r'([A-Za-z0-9]+) = (-?\d+(?:\.\d+)?)')

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
SECTION = PatternFill("solid", fgColor="D9E2F3")
BLUE = Font(name=FONT, color="0563C1", underline="single")

CLASS_ORDER = ["DRUID", "HUNTER", "MAGE", "PALADIN", "PRIEST", "ROGUE", "SHAMAN", "WARLOCK", "WARRIOR"]
REFERENCE = ("AttackPower", "SpellPower", "SpellDamage")


def readable(name):
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", name).replace("M P5", "MP5")


def read_weights(path):
    with open(path, encoding="utf-8") as handle:
        text = handle.read()
    info = INFO.search(text)
    specs = []
    for key, klass, name, source, stats, pseudo in SPEC.findall(text):
        specs.append({
            "key": key, "class": klass, "name": name, "source": source,
            "stats": {n: float(v) for n, v in PAIR.findall(stats)},
            "pseudo": {n: float(v) for n, v in PAIR.findall(pseudo)},
        })
    return specs, {"game": info.group(1), "generated": info.group(2), "specs": int(info.group(3)),
                   "repo": info.group(4), "commit": info.group(5)} if info else {}


def link(ws, row, column, text, url):
    cell = ws.cell(row=row, column=column, value=text)
    cell.hyperlink = url
    cell.font = BLUE
    return cell


def file_url(info, path):
    return "https://github.com/%s/blob/%s/%s" % (info["repo"], info["commit"], path)


def overview(wb, specs, info):
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "ForeverBuddy stat weights"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = ("What the addon scores gear with. Every number is read straight out of the WoWSims Classic "
                "simulators: each spec's sim.ts declares the weights its own sim solved for.")
    ws["A3"] = ("Source: %s at commit %s, read %s. These are CLASSIC weights, used until Forever's own exist."
                % (info["repo"], info["commit"], info["generated"]))
    ws["A4"] = ("Resistance stats are deliberately dropped. WoWSims' presets are written for raid fights that "
                "demand a resistance set, so nine specs carried a leftover fire resistance weight of 0.5 from "
                "Molten Core gearing. Left in, a fire resistance roll would read as an upgrade on a level 20 "
                "dungeon drop. Armour, defence, dodge and parry are kept: those are real for tanks at any level.")
    ws["A5"] = ("A weight is what one point of a stat is worth against 1 attack power or 1 spell damage, which "
                "are pinned at 1.00. Bigger is better. Blank means that sim gives the stat no value at all.")
    for row in range(2, 6):
        ws.cell(row=row, column=1).font = Font(name=FONT, italic=True, size=9)
    for i, (name, width) in enumerate(zip(
            ["Class", "Specs", "Spec keys for /fb spec", "Weights read from"], [14, 8, 46, 34]), start=1):
        c = ws.cell(row=7, column=i, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
        ws.column_dimensions[c.column_letter].width = width
    row = 8
    for klass in CLASS_ORDER:
        mine = [s for s in specs if s["class"] == klass]
        if not mine:
            continue
        ws.cell(row=row, column=1, value=klass.title()).font = Font(name=FONT)
        ws.cell(row=row, column=2, value=len(mine)).font = Font(name=FONT)
        ws.cell(row=row, column=3, value=", ".join(s["key"] for s in mine)).font = Font(name=FONT)
        folder = os.path.dirname(mine[0]["source"]) or "ui"
        link(ws, row, 4, folder.replace("ui/", "") + "/sim.ts", file_url(info, mine[0]["source"]))
        row += 1
    ws.cell(row=row + 1, column=1, value="Each class sheet links to the exact file behind every spec column.") \
        .font = Font(name=FONT, italic=True, size=9)
    ws.freeze_panes = "A8"


def class_sheet(wb, klass, specs, info):
    mine = sorted([s for s in specs if s["class"] == klass], key=lambda s: s["name"])
    ws = wb.create_sheet(klass.title())
    ws["A1"] = "%s stat weights" % klass.title()
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = ("Higher is better. Attack power and spell damage are the reference at 1.00. Blank means the sim "
                "gives that stat no value for this spec. Resistances are left out on purpose, see Overview.")
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    ws.column_dimensions["A"].width = 26

    head = ws.cell(row=4, column=1, value="Stat")
    head.font = Font(name=FONT, bold=True, color="FFFFFF")
    head.fill = HEAD
    for i, spec in enumerate(mine, start=2):
        c = ws.cell(row=4, column=i, value=spec["name"])
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
        ws.column_dimensions[c.column_letter].width = max(16, len(spec["name"]) + 2)
    ws.cell(row=5, column=1, value="Spec key (/fb spec)").font = Font(name=FONT, italic=True)
    ws.cell(row=6, column=1, value="Weights read from").font = Font(name=FONT, italic=True)
    for i, spec in enumerate(mine, start=2):
        ws.cell(row=6 - 1, column=i, value=spec["key"]).font = Font(name=FONT, italic=True)
        link(ws, 6, i, spec["source"].replace("ui/", ""), file_url(info, spec["source"]))

    row = 8
    ws.cell(row=row, column=1, value="Gear stats").fill = SECTION
    ws.cell(row=row, column=1).font = Font(name=FONT, bold=True)
    row += 1
    names = sorted({n for s in mine for n in s["stats"]},
                   key=lambda n: (n not in REFERENCE, readable(n)))
    for name in names:
        ws.cell(row=row, column=1, value=readable(name)).font = Font(name=FONT)
        for i, spec in enumerate(mine, start=2):
            if name in spec["stats"]:
                c = ws.cell(row=row, column=i, value=spec["stats"][name])
                c.font = Font(name=FONT)
                c.number_format = "0.00"
        row += 1

    pseudo = sorted({n for s in mine for n in s["pseudo"]})
    if pseudo:
        row += 1
        ws.cell(row=row, column=1, value="Weapon and speed").fill = SECTION
        ws.cell(row=row, column=1).font = Font(name=FONT, bold=True)
        row += 1
        for name in pseudo:
            ws.cell(row=row, column=1, value=readable(name)).font = Font(name=FONT)
            for i, spec in enumerate(mine, start=2):
                if name in spec["pseudo"]:
                    c = ws.cell(row=row, column=i, value=spec["pseudo"][name])
                    c.font = Font(name=FONT)
                    c.number_format = "0.00"
            row += 1
    ws.freeze_panes = "B7"


def main(argv=None):
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--weights", default=os.path.join(root, "Data", "Weights.lua"))
    parser.add_argument("--out", default=os.path.expanduser("~/Desktop/ForeverAddons/ForeverBuddy-stat-weights.xlsx"))
    args = parser.parse_args(argv)
    specs, info = read_weights(args.weights)
    if not specs or not info:
        sys.exit("could not read %s" % args.weights)
    wb = Workbook()
    overview(wb, specs, info)
    for klass in CLASS_ORDER:
        if any(s["class"] == klass for s in specs):
            class_sheet(wb, klass, specs, info)
    wb.save(args.out)
    print("%d specs across %d sheets written to %s" % (len(specs), len(wb.sheetnames) - 1, args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
